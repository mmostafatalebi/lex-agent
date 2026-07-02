"""Regenerate the sample MSA fixtures (PDF + DOCX) from the Markdown source.

Idempotent and re-runnable. Reads ``fixtures/sample_msa.md`` and writes
``fixtures/sample_msa.pdf`` and ``fixtures/sample_msa.docx`` alongside it.

Usage::

    uv run python scripts/build_fixtures.py
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "fixtures"
SOURCE = FIXTURES_DIR / "sample_msa.md"
PDF_OUT = FIXTURES_DIR / "sample_msa.pdf"
DOCX_OUT = FIXTURES_DIR / "sample_msa.docx"


@dataclass(frozen=True)
class Block:
    """A parsed Markdown block: a title, a heading, or a body paragraph."""

    kind: str  # "title" | "heading" | "body"
    text: str


def parse_markdown(md: str) -> list[Block]:
    """Parse the controlled fixture Markdown into an ordered list of blocks."""
    blocks: list[Block] = []
    for chunk in md.split("\n\n"):
        line = chunk.strip()
        if not line:
            continue
        if line.startswith("## "):
            blocks.append(Block("heading", line[3:].strip()))
        elif line.startswith("# "):
            blocks.append(Block("title", line[2:].strip()))
        else:
            # Collapse hard-wrapped lines within a paragraph into one line.
            body = " ".join(part.strip() for part in line.splitlines())
            blocks.append(Block("body", body))
    return blocks


def build_pdf(blocks: list[Block], out: Path) -> None:
    """Render the blocks to a clean serif PDF on letter-size pages."""
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "MSATitle",
        parent=styles["Title"],
        fontName="Times-Bold",
        fontSize=16,
        spaceAfter=18,
    )
    heading_style = ParagraphStyle(
        "MSAHeading",
        parent=styles["Heading2"],
        fontName="Times-Bold",
        fontSize=12,
        spaceBefore=14,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "MSABody",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=11,
        leading=15,
        alignment=TA_JUSTIFY,
        spaceAfter=8,
    )

    doc = SimpleDocTemplate(
        str(out),
        pagesize=LETTER,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
        leftMargin=1 * inch,
        rightMargin=1 * inch,
        title="Master Service Agreement",
    )
    flow: list[object] = []
    for block in blocks:
        if block.kind == "title":
            flow.append(Paragraph(block.text, title_style))
        elif block.kind == "heading":
            flow.append(Paragraph(block.text, heading_style))
        else:
            flow.append(Paragraph(block.text, body_style))
    flow.append(Spacer(1, 12))
    doc.build(flow)


def build_docx(blocks: list[Block], out: Path) -> None:
    """Render the blocks to a DOCX with matching structure."""
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    for block in blocks:
        if block.kind == "title":
            document.add_heading(block.text, level=0)
        elif block.kind == "heading":
            document.add_heading(block.text, level=1)
        else:
            document.add_paragraph(block.text)
    document.save(str(out))


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source markdown: {SOURCE}")
    blocks = parse_markdown(SOURCE.read_text(encoding="utf-8"))
    build_pdf(blocks, PDF_OUT)
    build_docx(blocks, DOCX_OUT)
    print("Wrote fixtures:")
    print(f"  {PDF_OUT}")
    print(f"  {DOCX_OUT}")


if __name__ == "__main__":
    main()
